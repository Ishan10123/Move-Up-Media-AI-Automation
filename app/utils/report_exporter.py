import os
import re

from datetime import datetime

from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.enum.text import (
    WD_PARAGRAPH_ALIGNMENT
)

from fpdf import FPDF


def ensure_reports_directory():

    os.makedirs(
        "reports",
        exist_ok=True
    )


def clean_text(text):

    replacements = {
        "–": "-",
        "—": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "•": "-",
        "…": "...",
        "\u00a0": " ",
        "\t": " "
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    cleaned = text.encode(
        "latin-1",
        "ignore"
    ).decode("latin-1")

    return cleaned.strip()


def is_main_heading(text):

    patterns = [
        r"^\d+\.",
        r"^#+",
        r"^Executive Summary",
        r"^Strategic",
        r"^Engagement",
        r"^Content",
        r"^Weekly"
    ]

    for pattern in patterns:

        if re.match(
            pattern,
            text,
            re.IGNORECASE
        ):

            return True

    return False


def is_sub_point(text):

    return (
        text.startswith("-")
        or text.startswith("•")
    )


def normalize_line(line):

    line = clean_text(line)

    line = line.replace(
        "##",
        ""
    )

    line = line.replace(
        "#",
        ""
    )

    return line.strip()


def add_docx_cover_page(
    doc,
    channel_name
):

    title = doc.add_heading(
        f"{channel_name} AI Performance Report",
        level=0
    )

    title.alignment = (
        WD_PARAGRAPH_ALIGNMENT.CENTER
    )

    title.style.font.size = Pt(24)

    subtitle = doc.add_paragraph(
        "MoveUp Media - AI Content Operations Intelligence"
    )

    subtitle.alignment = (
        WD_PARAGRAPH_ALIGNMENT.CENTER
    )

    subtitle.style.font.size = Pt(12)

    timestamp = doc.add_paragraph(
        f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    timestamp.alignment = (
        WD_PARAGRAPH_ALIGNMENT.CENTER
    )

    timestamp.style.font.size = Pt(10)

    doc.add_page_break()


def export_to_docx(
    report,
    filename,
    channel_name="Unknown"
):

    ensure_reports_directory()

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    add_docx_cover_page(
        doc,
        channel_name
    )

    lines = report.split("\n")

    for raw_line in lines:

        line = normalize_line(
            raw_line
        )

        if not line:

            continue

        if is_main_heading(line):

            heading = doc.add_heading(
                line,
                level=1
            )

            heading.style.font.size = Pt(16)

        elif is_sub_point(line):

            bullet_text = (
                line.replace(
                    "-",
                    ""
                ).replace(
                    "•",
                    ""
                ).strip()
            )

            bullet = doc.add_paragraph(
                bullet_text,
                style="List Bullet"
            )

            bullet.style.font.size = Pt(11)

        else:

            paragraph = doc.add_paragraph(
                line
            )

            paragraph.style.font.size = Pt(11)

    doc.save(filename)

    return filename


class EnterprisePDF(FPDF):

    def header(self):

        self.set_font(
            "Arial",
            "B",
            20
        )

        self.cell(
            0,
            12,
            self.report_title,
            ln=True,
            align="C"
        )

        self.set_font(
            "Arial",
            "",
            11
        )

        self.cell(
            0,
            8,
            "MoveUp Media - AI Content Operations Intelligence",
            ln=True,
            align="C"
        )

        self.set_font(
            "Arial",
            "I",
            9
        )

        self.cell(
            0,
            6,
            f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            ln=True,
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Arial",
            "I",
            9
        )

        self.cell(
            0,
            10,
            f"Page {self.page_no()}",
            align="C"
        )


def add_pdf_heading(
    pdf,
    text
):

    pdf.ln(4)

    pdf.set_font(
        "Arial",
        "B",
        15
    )

    pdf.multi_cell(
        180,
        10,
        text
    )

    pdf.ln(2)


def add_pdf_paragraph(
    pdf,
    text
):

    pdf.set_font(
        "Arial",
        "",
        11
    )

    pdf.multi_cell(
        180,
        7,
        text
    )

    pdf.ln(1)


def add_pdf_bullet(
    pdf,
    text
):

    bullet_text = (
        "- "
        +
        text.replace(
            "-",
            ""
        ).replace(
            "•",
            ""
        ).strip()
    )

    pdf.set_font(
        "Arial",
        "",
        11
    )

    pdf.multi_cell(
        180,
        7,
        bullet_text
    )


def export_to_pdf(
    report,
    filename,
    channel_name="Unknown"
):

    ensure_reports_directory()

    pdf = EnterprisePDF()

    pdf.report_title = (
        f"{channel_name} AI Performance Report"
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )

    pdf.set_left_margin(15)
    pdf.set_right_margin(15)

    pdf.add_page()

    lines = report.split("\n")

    for raw_line in lines:

        try:

            line = normalize_line(
                raw_line
            )

            if not line:

                pdf.ln(3)

                continue

            if len(line.strip()) <= 1:

                continue

            if is_main_heading(line):

                add_pdf_heading(
                    pdf,
                    line
                )

            elif is_sub_point(line):

                add_pdf_bullet(
                    pdf,
                    line
                )

            else:

                add_pdf_paragraph(
                    pdf,
                    line
                )

        except Exception:

            continue

    pdf.output(filename)

    return filename